"""Render the chief intelligence officer's plaintext briefing as a docx.

Parses the three top-level sections (MARKET OVERVIEW, COMPANY-SPECIFIC
LEADS, IMMEDIATE PRIORITIES) from the briefing schema defined in
agents/chief_intelligence_officer.py and writes a Word file into the
per-run output directory provided via `state["output_dir"]`.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt, RGBColor

from workflows.lead_gen_fcc_intel.state import AgentState

_DEFAULT_OUTPUT_DIR = Path("output")
_DOCX_FILENAME = "fcc_intel_briefing.docx"
_SECTION_HEADERS = ("MARKET OVERVIEW", "COMPANY-SPECIFIC LEADS", "IMMEDIATE PRIORITIES")
_HEADER_NAVY = RGBColor(0x0B, 0x2A, 0x4A)
_BULLET_STYLE = "List Bullet"


def _split_sections(report: str) -> dict[str, str]:
    """Split the briefing on its three top-level section headers."""
    pattern = re.compile(
        r"^(?P<header>" + "|".join(re.escape(h) for h in _SECTION_HEADERS) + r")\s*$",
        re.MULTILINE,
    )
    matches = list(pattern.finditer(report))
    sections: dict[str, str] = {h: "" for h in _SECTION_HEADERS}
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(report)
        sections[m.group("header")] = report[m.end():end].strip()
    return sections


def _add_h1(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = _HEADER_NAVY


def _add_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = _HEADER_NAVY


def _add_h3(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)


def _add_meta(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    label_run = p.add_run(f"{label}: ")
    label_run.bold = True
    p.add_run(value)


def _render_bullet_block(doc: Document, block: str) -> None:
    """Render a hyphen-bulleted block, preserving 2-space indent as nested."""
    for line in block.splitlines():
        if not line.strip():
            continue
        stripped = line.lstrip()
        indent = len(line) - len(stripped)
        if stripped.startswith("- "):
            text = stripped[2:].strip()
            level = 1 if indent >= 2 else 0
            p = doc.add_paragraph(text, style=_BULLET_STYLE)
            p.paragraph_format.left_indent = Pt(18 * (level + 1))
        else:
            doc.add_paragraph(line)


def _render_market_overview(doc: Document, body: str) -> None:
    _add_h2(doc, "Market Overview")
    _render_bullet_block(doc, body)


def _render_company_leads(doc: Document, body: str) -> None:
    _add_h2(doc, "Company-Specific Leads")
    # Companies are separated by blank lines; the first line of each block is
    # the company name in CAPS, followed by hyphen-bulleted attributes.
    blocks = [b for b in re.split(r"\n\s*\n", body.strip()) if b.strip()]
    for block in blocks:
        lines = block.splitlines()
        company_name = lines[0].strip()
        _add_h3(doc, company_name)
        _render_bullet_block(doc, "\n".join(lines[1:]))


def _render_immediate_priorities(doc: Document, body: str) -> None:
    _add_h2(doc, "Immediate Priorities")
    _render_bullet_block(doc, body)


def _build_docx(report: str, generated_at: dt.datetime) -> Document:
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    title_run = title.add_run("FCC Lead Generation Intelligence Briefing")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = _HEADER_NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    subtitle.add_run(generated_at.strftime("%A, %d %B %Y · %H:%M")).italic = True

    doc.add_paragraph()

    sections = _split_sections(report)
    _render_market_overview(doc, sections["MARKET OVERVIEW"])
    _render_company_leads(doc, sections["COMPANY-SPECIFIC LEADS"])
    _render_immediate_priorities(doc, sections["IMMEDIATE PRIORITIES"])

    return doc


def format_report_node(state: AgentState) -> AgentState:
    report = state.get("chief_intelligence_report")
    if not report:
        return state

    out_dir = Path(state.get("output_dir") or _DEFAULT_OUTPUT_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / _DOCX_FILENAME

    doc = _build_docx(report, dt.datetime.now())
    doc.save(out_path)

    state["formatted_report_path"] = str(out_path)
    return state
